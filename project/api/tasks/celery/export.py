# services/socamas/project/api/tasks/celery/export.py
import time
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..base import socweb
from project import create_app
from project.api.utils.time import ts2locale
from project.config import DevelopmentConfig
from project.api.bizs import DomainArchivedBiz, WebsiteArchivedBiz


@socweb.task(queue='export')
def export_data(payload):
    table = payload.get('table', 'website')

    app = create_app()
    app.app_context().push()

    if table == 'domain':
        biz = DomainArchivedBiz()
        shown_fields = ['name', 'sponsor', 'sponsor_type', 'icp_number', 'city_code', 'website_count']  # 主域名导出字段
    else:
        biz = WebsiteArchivedBiz()
        shown_fields = ['url', 'title', 'ip', 'city_code', 'region_code', 'host_dept', 'host_type', 'industries',
                        'web_type', 'http_status', 'http_status_list']  # 子域名导出字段
    data = biz.query(**payload, fields=shown_fields)
    try:
        filename = generate_file(table, data, shown_fields)
        print(f"导出完成，文件名：{filename}")
    except Exception as e:
        print(f"生成文件异常，{e}")
        return
    return filename


def generate_file(table, data, fields=None):
    # 生成xlsx文件，保存到本地
    filename = f'project/api/data/{table}.xlsx'
    wb = load_workbook(filename=filename)
    ws = wb.active

    filename = "%s.xlsx" % int(time.time())
    full_path = '{}/{}'.format(DevelopmentConfig.EXPORT_PATH, filename)

    for record in data['records']:
        row_data = []
        for field in fields:
            if field == 'id':
                continue
            if field in ['create_time', 'update_time']:
                row_data.append(ts2locale(record[field]))
            elif field in ['icp_updated', 'website_count_updated']:
                value = 0 if not record.get(field) else record[field]
                day_count = int((time.time() - value) / (60 * 60 * 24))
                row_data.append('{}天'.format(day_count))
            elif field in ['tags', 'industries']:
                value = record[field] if record[field] else []
                row_data.append(','.join(value))
            else:
                row_data.append(str(record[field]).strip())
        ws.append(row_data)

    displayName = '子域名列表' if table == 'website' else '主域名列表'
    tab = Table(displayName=displayName, ref="A1:F{}".format(len(data['records']) + 1))
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True,
                           showColumnStripes=True)

    tab.tableStyleInfo = style
    ws.add_table(tab)

    wb.save(full_path)
    return filename

